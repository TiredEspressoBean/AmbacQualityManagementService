"""The master migration workbook (2026-10-01): one workbook, a sheet per table, loaded
in order in one transaction — dry run first, all or nothing, re-uploadable.
See Tracker.services.core.master_workbook."""
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase

from Tracker.utils.tenant_context import reset_current_tenant, set_current_tenant_id

User = get_user_model()


class MasterWorkbookTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        from Tracker.models import Tenant
        cls.tenant = Tenant.objects.create(name="Migration", slug="master-workbook-1001")
        token = set_current_tenant_id(cls.tenant.id)
        try:
            cls.admin = User.objects.create_user(username="mw-admin", email="mw-admin@x.test",
                                                 password="x", tenant=cls.tenant, is_staff=True)
            cls.admin.is_superuser = True
            cls.admin.save(update_fields=["is_superuser"])
        finally:
            reset_current_tenant(token)

    def setUp(self):
        self.client.force_authenticate(user=self.admin)
        self.client.credentials(HTTP_X_TENANT_ID=str(self.tenant.id))
        self._token = set_current_tenant_id(self.tenant.id)

    def tearDown(self):
        reset_current_tenant(self._token)

    # -- helpers ---------------------------------------------------------------

    def _blank(self, client=None):
        resp = (client or self.client).get("/api/MasterWorkbook/template/")
        self.assertEqual(resp.status_code, 200)
        return load_workbook(io.BytesIO(resp.content))

    @staticmethod
    def _put(wb, title, *rows):
        """Append rows to a sheet, by header (the `*` ignored)."""
        ws = wb[title]
        heads = [str(c.value).rstrip("*") for c in ws[1]]
        for row in rows:
            unknown = set(row) - set(heads)
            assert not unknown, f"{title} has no column {unknown}"
            ws.append([row.get(h) for h in heads])

    def _run(self, wb, dry_run=True, client=None):
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "migration.xlsx"
        return (client or self.client).post("/api/MasterWorkbook/run/", {
            "file": buf, "dry_run": "true" if dry_run else "false"}, format="multipart")

    def _sheet(self, body, title):
        return next(s for s in body["sheets"] if s["sheet"] == title)

    def _workbook(self):
        """A small plant: a supplier, a part, its stock, a person and their training."""
        wb = self._blank()
        self._put(wb, "Users", {"Email": "kim@x.test", "First Name": "Kim", "Last Name": "Lee"})
        self._put(wb, "Companies", {"name": "Acme", "description": "Seals",
                                    "is_customer": "FALSE", "is_supplier": "TRUE"})
        self._put(wb, "Part Types", {"name": "Seal", "ERP_id": "SEAL-1", "can_buy": "TRUE"})
        self._put(wb, "Training Types", {"name": "Torque", "description": "Torque wrench use"})
        self._put(wb, "Training Records", {"user": "kim@x.test", "training_type": "Torque",
                                           "completed_date": "2026-09-01", "level": "3"})
        self._put(wb, "Stock on Hand", {"Lot Number": "L-100", "Item": "SEAL-1",
                                        "Quantity": "40", "Supplier": "Acme", "Location": "Rack 1"})
        return wb

    # -- the workbook ------------------------------------------------------------

    def test_the_blank_workbook_has_a_sheet_per_table_in_load_order(self):
        from Tracker.services.core.master_workbook import SHEETS
        wb = self._blank()
        self.assertEqual(wb.sheetnames, ["Read me"] + [s.title for s in SHEETS])
        self.assertIn("Lot Number*", [c.value for c in wb["Stock on Hand"][1]])
        # Earlier sheets come first: a later row names what an earlier one made.
        order = wb.sheetnames
        for earlier, later in [("Companies", "Part Types"), ("Part Types", "BOMs"),
                               ("Users", "Training Records"), ("Work Orders", "Parts")]:
            self.assertLess(order.index(earlier), order.index(later))

    def test_a_dry_run_reports_everything_and_keeps_nothing(self):
        from Tracker.models import Companies, MaterialLot, TrainingRecord
        resp = self._run(self._workbook())
        self.assertEqual(resp.status_code, 200, resp.content)
        body = resp.json()
        self.assertTrue(body["dry_run"])
        self.assertFalse(body["loaded"])
        self.assertEqual(body["totals"]["errors"], 0, body)
        self.assertEqual(body["totals"]["created"], 6)
        self.assertFalse(Companies.objects.filter(name="Acme").exists())
        self.assertFalse(MaterialLot.objects.filter(lot_number="L-100").exists())
        self.assertFalse(TrainingRecord.objects.exists())
        self.assertFalse(User.objects.filter(email="kim@x.test").exists())

    def test_a_load_keeps_it_and_later_sheets_name_earlier_ones(self):
        from Tracker.models import MaterialLot, TrainingRecord
        resp = self._run(self._workbook(), dry_run=False)
        body = resp.json()
        self.assertTrue(body["loaded"], body)
        lot = MaterialLot.objects.get(lot_number="L-100")
        self.assertEqual((lot.status, lot.quantity, lot.quantity_remaining, lot.storage_location),  # created from the sheet
                         ("ACCEPTED", Decimal("40"), Decimal("40"), "Rack 1"))
        self.assertEqual((lot.supplier.name, lot.material_type.ERP_id), ("Acme", "SEAL-1"))
        record = TrainingRecord.objects.get()
        self.assertEqual((record.user.email, record.training_type.name, record.level),
                         ("kim@x.test", "Torque", 3))

    def test_one_bad_row_loads_nothing(self):
        from Tracker.models import Companies
        wb = self._workbook()
        self._put(wb, "Stock on Hand", {"Lot Number": "L-101", "Item": "NO-SUCH-PART",
                                        "Quantity": "1"})
        body = self._run(wb, dry_run=False).json()
        self.assertFalse(body["loaded"])
        stock = self._sheet(body, "Stock on Hand")
        self.assertEqual(stock["errors"], 1)
        self.assertEqual(stock["rows"][0]["row"], 3)  # the spreadsheet row: header is 1
        self.assertIn("NO-SUCH-PART", stock["rows"][0]["detail"])
        self.assertFalse(Companies.objects.filter(name="Acme").exists())

    def test_uploading_it_again_updates_rather_than_duplicates(self):
        from Tracker.models import Companies, MaterialLot, TrainingRecord
        self.assertTrue(self._run(self._workbook(), dry_run=False).json()["loaded"])
        wb = self._workbook()
        wb["Stock on Hand"]["C2"] = "45"  # recounted
        body = self._run(wb, dry_run=False).json()
        self.assertTrue(body["loaded"], body)
        self.assertEqual(body["totals"]["created"], 0, body)
        self.assertEqual(Companies.objects.filter(name="Acme", archived=False,
                                                  is_current_version=True).count(), 1)
        self.assertEqual(MaterialLot.objects.get(lot_number="L-100").quantity_remaining, Decimal("45"))
        self.assertEqual(TrainingRecord.objects.count(), 1)

    def test_stock_already_in_use_is_left_alone(self):
        from Tracker.models import MaterialLot
        self.assertTrue(self._run(self._workbook(), dry_run=False).json()["loaded"])
        MaterialLot.objects.filter(lot_number="L-100").update(quantity_remaining=Decimal("30"))
        wb = self._workbook()
        wb["Stock on Hand"]["C2"] = "45"
        body = self._run(wb, dry_run=False).json()
        self.assertTrue(body["loaded"], body)
        self.assertIn("left alone", self._sheet(body, "Stock on Hand")["rows"][0]["detail"])
        lot = MaterialLot.objects.get(lot_number="L-100")
        self.assertEqual((lot.quantity, lot.quantity_remaining), (Decimal("40"), Decimal("30")))

    def test_a_training_record_is_added_never_rewritten(self):
        self.assertTrue(self._run(self._workbook(), dry_run=False).json()["loaded"])
        wb = self._workbook()
        wb["Training Records"]["D2"] = "4"  # a different level on the same record
        body = self._run(wb).json()
        records = self._sheet(body, "Training Records")
        self.assertEqual(records["errors"], 1)
        self.assertIn("evidence", records["rows"][0]["detail"])

    def test_owner_must_be_a_customer(self):
        wb = self._workbook()
        self._put(wb, "Stock on Hand", {"Lot Number": "L-102", "Item": "SEAL-1",
                                        "Quantity": "2", "Owner": "Acme"})
        stock = self._sheet(self._run(wb).json(), "Stock on Hand")
        self.assertEqual(stock["errors"], 1)
        self.assertIn("customer", stock["rows"][0]["detail"])

    def test_a_sheet_needs_its_tables_permissions(self):
        from django.contrib.auth.models import Permission
        from Tracker.models import TenantGroup, UserRole
        clerk = User.objects.create_user(username="mw-clerk", email="mw-clerk@x.test",
                                         password="x", tenant=self.tenant)
        group = TenantGroup.objects.create(tenant=self.tenant, name="MW Clerk", is_custom=True)
        group.permissions.add(*Permission.objects.filter(
            codename__in=["add_companies", "change_companies", "view_companies"]))
        UserRole.objects.create(user=clerk, group=group)
        clerk.clear_permission_cache(self.tenant)
        client = APIClient()
        client.force_authenticate(user=clerk)
        client.credentials(HTTP_X_TENANT_ID=str(self.tenant.id))

        allowed = {s["title"]: s["allowed"] for s in client.get("/api/MasterWorkbook/sheets/").json()}
        self.assertTrue(allowed["Companies"])
        self.assertFalse(allowed["Stock on Hand"])
        body = self._run(self._workbook(), client=client).json()
        self.assertEqual(self._sheet(body, "Companies")["errors"], 0)
        self.assertIn("permission", self._sheet(body, "Stock on Hand")["detail"])

    def test_sheets_that_arent_the_workbooks_are_reported_not_loaded(self):
        wb = self._workbook()
        wb.create_sheet("My notes").append(["anything"])
        body = self._run(wb).json()
        self.assertEqual(body["ignored_sheets"], ["My notes"])

    def test_an_empty_workbook_is_refused(self):
        resp = self._run(self._blank())
        self.assertEqual(resp.status_code, 400)

    # -- 2026-10-01, second pass ---------------------------------------------------

    def test_users_load_without_invitations(self):
        from Tracker.models import UserInvitation
        self.assertTrue(self._run(self._workbook(), dry_run=False).json()["loaded"])
        kim = User.objects.get(email="kim@x.test")
        self.assertFalse(UserInvitation.objects.filter(user=kim).exists())

    def test_cores_load_as_received_with_their_part(self):
        from Tracker.models import Core
        wb = self._workbook()
        self._put(wb, "Companies", {"name": "Fleet Co", "description": "Customer",
                                    "is_customer": "TRUE", "is_supplier": "FALSE"})
        self._put(wb, "Part Types", {"name": "Injector Core", "ERP_id": "INJ-C"})
        self._put(wb, "Cores", {"core_type": "Injector Core", "customer": "Fleet Co",
                                "serial_number": "SN-1", "received_date": "2026-09-01",
                                "source_type": "CUSTOMER_RETURN", "condition_grade": "B"})
        body = self._run(wb, dry_run=False).json()
        self.assertTrue(body["loaded"], body)
        core = Core.objects.get(serial_number="SN-1")
        self.assertEqual((core.status, core.customer.name), ("RECEIVED", "Fleet Co"))
        self.assertIsNotNone(core.part_id)

    def test_a_core_from_a_supplier_only_company_is_refused(self):
        wb = self._workbook()
        self._put(wb, "Part Types", {"name": "Injector Core", "ERP_id": "INJ-C"})
        self._put(wb, "Cores", {"core_type": "Injector Core", "customer": "Acme",
                                "received_date": "2026-09-01", "source_type": "CUSTOMER_RETURN",
                                "condition_grade": "B"})
        cores = self._sheet(self._run(wb).json(), "Cores")
        self.assertEqual(cores["errors"], 1, cores)
        self.assertIn("customer", cores["rows"][0]["detail"])

    def test_a_filled_workbook_uploads_back_unchanged(self):
        self.assertTrue(self._run(self._workbook(), dry_run=False).json()["loaded"])
        resp = self.client.get("/api/MasterWorkbook/template/?filled=true")
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb["Stock on Hand"]["A2"].value, "L-100")
        self.assertEqual(wb["Users"].max_row, 3)  # header, the admin, Kim
        body = self._run(wb, dry_run=False).json()
        self.assertTrue(body["loaded"], body)
        self.assertEqual(body["totals"]["created"], 0, body)
        # Matched and unchanged: "no change", not an update; and no warnings for the
        # export's read-only columns (they aren't written into the filled workbook).
        self.assertEqual(body["totals"]["updated"], 0, body)
        self.assertFalse([r for s in body["sheets"] for r in s["rows"]], body)

    def test_loaded_users_are_invited_together_at_go_live(self):
        from Tracker.models import UserInvitation
        from Tracker.services.core.tenant_membership import suspend_membership
        wb = self._workbook()
        self._put(wb, "Users", {"Email": "lee@x.test", "First Name": "Lee"})
        self.assertTrue(self._run(wb, dry_run=False).json()["loaded"])
        # Removed before go-live: never invited by the bulk action.
        suspend_membership(User.objects.get(email="lee@x.test"), self.tenant, by=self.admin)
        listed = self.client.get("/api/User/uninvited/").json()
        self.assertEqual(listed["emails"], ["kim@x.test"])
        resp = self.client.post("/api/User/invite-uninvited/")
        self.assertEqual(resp.json()["invited"], 1, resp.content)
        self.assertTrue(UserInvitation.objects.filter(user__email="kim@x.test").exists())
        self.assertFalse(UserInvitation.objects.filter(user__email="lee@x.test").exists())
        self.assertEqual(self.client.get("/api/User/uninvited/").json()["count"], 0)
