"""
Tests for CSV Import/Export functionality.

Tests:
- CSV/Excel file parsing with encoding detection
- Template generation
- Import serializers with FK resolution
- ViewSet mixins (import/export endpoints)
"""

import io
import csv
from unittest.mock import patch, MagicMock
from uuid import uuid4

from django.test import TestCase, override_settings
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from rest_framework import status

from Tracker.tests.base import TenantContextMixin

from Tracker.services.csv_utils import (
    detect_encoding,
    decode_file_content,
    normalize_header,
    create_field_mapping,
    parse_csv_file,
    parse_date,
    parse_boolean,
    parse_integer,
    parse_decimal,
    ImportResult,
)
from Tracker.services.template_generator import (
    TemplateField,
    TemplateGenerator,
    template_from_model,
    introspect_model,
    get_part_types_template,
)
from Tracker.serializers.csv_import import (
    ImportMode,
    BaseCSVImportSerializer,
    PartTypesCSVImportSerializer,
    PartsCSVImportSerializer,
    OrdersCSVImportSerializer,
    get_csv_import_serializer,
    get_or_create_import_serializer,
    create_import_serializer_for_model,
)

User = get_user_model()


class CsvUtilsTests(TestCase):
    """Tests for csv_utils.py functions."""

    def test_detect_encoding_utf8(self):
        """Test UTF-8 encoding detection."""
        content = "Hello, World!".encode("utf-8")
        encoding = detect_encoding(content)
        self.assertIn(encoding.lower(), ["utf-8", "ascii"])

    def test_detect_encoding_latin1(self):
        """Test Latin-1 encoding detection."""
        # Content with Latin-1 specific characters
        content = "Café résumé naïve".encode("latin-1")
        encoding = detect_encoding(content)
        # Should detect some encoding (may vary)
        self.assertIsNotNone(encoding)

    def test_decode_file_content_utf8(self):
        """Test decoding UTF-8 content."""
        original = "Hello, World! こんにちは"
        content = original.encode("utf-8")
        decoded = decode_file_content(content)
        self.assertEqual(decoded, original)

    def test_decode_file_content_with_fallback(self):
        """Test decoding with fallback when encoding fails."""
        content = b"Hello, World!"
        decoded = decode_file_content(content, encoding="invalid-encoding")
        self.assertIn("Hello", decoded)

    def test_normalize_header(self):
        """Test header normalization."""
        self.assertEqual(normalize_header("  Part Type  "), "part_type")
        self.assertEqual(normalize_header("WO Number"), "wo_number")
        self.assertEqual(normalize_header("Item's Name"), "item's_name")
        self.assertEqual(normalize_header("ERP-ID"), "erp_id")

    def test_create_field_mapping(self):
        """Test field mapping creation."""
        headers = ["Part Type", "WO Number", "Quantity"]
        field_map = {
            "part_type": "part_type_id",
            "wo_number": "ERP_id",
        }
        mapping = create_field_mapping(headers, field_map)
        self.assertEqual(mapping["Part Type"], "part_type_id")
        self.assertEqual(mapping["WO Number"], "ERP_id")
        self.assertEqual(mapping["Quantity"], "quantity")

    def test_parse_csv_file(self):
        """Test CSV file parsing."""
        csv_content = "Name,Quantity,Status\nWidget,100,active\nGadget,50,pending"
        file = io.BytesIO(csv_content.encode("utf-8"))
        rows, headers = parse_csv_file(file)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["name"], "Widget")
        self.assertEqual(rows[0]["quantity"], "100")
        self.assertEqual(rows[1]["name"], "Gadget")

    def test_parse_csv_file_with_field_mapping(self):
        """Test CSV parsing with field mapping."""
        csv_content = "Part Type,WO Number\nWidget,WO-001"
        file = io.BytesIO(csv_content.encode("utf-8"))
        field_map = {"part_type": "part_type_id", "wo_number": "ERP_id"}
        rows, headers = parse_csv_file(file, field_map=field_map)

        self.assertEqual(rows[0]["part_type_id"], "Widget")
        self.assertEqual(rows[0]["ERP_id"], "WO-001")

    def test_parse_date_various_formats(self):
        """Test date parsing with various formats."""
        from datetime import datetime

        # YYYY-MM-DD
        result = parse_date("2024-01-15")
        self.assertEqual(result.year, 2024)
        self.assertEqual(result.month, 1)
        self.assertEqual(result.day, 15)

        # MM/DD/YYYY
        result = parse_date("01/15/2024")
        self.assertEqual(result.year, 2024)
        self.assertEqual(result.month, 1)

        # Empty/None
        self.assertIsNone(parse_date(None))
        self.assertIsNone(parse_date(""))

    def test_parse_boolean(self):
        """Test boolean parsing."""
        self.assertTrue(parse_boolean("true"))
        self.assertTrue(parse_boolean("yes"))
        self.assertTrue(parse_boolean("1"))
        self.assertTrue(parse_boolean(True))

        self.assertFalse(parse_boolean("false"))
        self.assertFalse(parse_boolean("no"))
        self.assertFalse(parse_boolean("0"))
        self.assertFalse(parse_boolean(False))

        self.assertIsNone(parse_boolean(None))

    def test_parse_integer(self):
        """Test integer parsing."""
        self.assertEqual(parse_integer("100"), 100)
        self.assertEqual(parse_integer(100), 100)
        self.assertEqual(parse_integer(100.5), 100)
        self.assertEqual(parse_integer("100.5"), 100)
        self.assertIsNone(parse_integer(None))
        self.assertIsNone(parse_integer(""))

    def test_parse_decimal(self):
        """Test decimal parsing."""
        self.assertEqual(parse_decimal("100.5"), 100.5)
        self.assertEqual(parse_decimal(100), 100.0)
        self.assertIsNone(parse_decimal(None))
        self.assertIsNone(parse_decimal(""))

    def test_import_result(self):
        """Test ImportResult tracking."""
        result = ImportResult()

        result.add_created(1, "uuid-1")
        result.add_updated(2, "uuid-2", warnings=["Warning 1"])
        result.add_error(3, {"field": "Error message"})

        self.assertEqual(result.total, 3)
        self.assertEqual(result.created, 1)
        self.assertEqual(result.updated, 1)
        self.assertEqual(result.errors, 1)

        response = result.to_response()
        self.assertEqual(response["summary"]["total"], 3)
        self.assertEqual(len(response["results"]), 3)


class TemplateGeneratorTests(TestCase):
    """Tests for template_generator.py."""

    def test_template_field_header_name(self):
        """Test required field marker in header."""
        field = TemplateField("name", required=True)
        self.assertEqual(field.header_name, "name*")

        field2 = TemplateField("notes", required=False)
        self.assertEqual(field2.header_name, "notes")

    def test_template_field_hint_text(self):
        """Test hint text generation."""
        field = TemplateField(
            "status",
            description="Current status",
            choices=["active", "inactive"],
        )
        hint = field.hint_text
        self.assertIn("Current status", hint)
        self.assertIn("active", hint)

    def test_generate_csv_template(self):
        """Test CSV template generation."""
        generator = TemplateGenerator(
            model_name="TestModel",
            fields=[
                TemplateField("name", required=True, example="Test Name"),
                TemplateField("quantity", example=100),
            ],
        )
        content = generator.generate_csv()

        # Parse the CSV to verify
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        rows = list(reader)

        self.assertEqual(rows[0], ["name*", "quantity"])
        # Header only: a sample row under it was imported as data.
        self.assertEqual(len(rows), 1)

    def test_get_part_types_template(self):
        """Test pre-built part types template."""
        template = get_part_types_template()
        self.assertEqual(template.model_name, "PartTypes")
        self.assertTrue(any(f.name == "name" and f.required for f in template.fields))


class CsvImportSerializerTests(TenantContextMixin, TestCase):
    """Tests for csv_import.py serializers."""

    @classmethod
    def setUpTestData(cls):
        """Create test data."""
        from Tracker.models import Tenant
        cls.tenant = Tenant.objects.create(name="CSV Test Tenant", slug="csv-test")
        cls.set_tenant_context_class(cls.tenant)
        cls.user = User.objects.create_user(
            username="testuser",
            email="test@example.com",
            password="testpass123",
            tenant=cls.tenant,
        )

    def test_get_csv_import_serializer(self):
        """Test serializer lookup by model name."""
        self.assertEqual(
            get_csv_import_serializer("part-types"),
            PartTypesCSVImportSerializer
        )
        self.assertEqual(
            get_csv_import_serializer("parts"),
            PartsCSVImportSerializer
        )
        self.assertEqual(
            get_csv_import_serializer("orders"),
            OrdersCSVImportSerializer
        )
        self.assertIsNone(get_csv_import_serializer("unknown"))

    def test_import_mode_constants(self):
        """Test import mode constants."""
        self.assertEqual(ImportMode.CREATE, "create")
        self.assertEqual(ImportMode.UPDATE, "update")
        self.assertEqual(ImportMode.UPSERT, "upsert")

    def test_part_types_serializer_create(self):
        """Test PartTypes CSV import - create mode."""
        serializer = PartTypesCSVImportSerializer(
            user=self.user,
            mode=ImportMode.CREATE,
        )

        row_data = {
            "name": "Test Part Type",
            "ERP_id": "PT-001",
            "ID_prefix": "TST-",
        }

        instance, created, warnings = serializer.import_row(row_data)

        self.assertTrue(created)
        self.assertEqual(instance.name, "Test Part Type")
        self.assertEqual(instance.ERP_id, "PT-001")

    def test_part_types_serializer_upsert_existing(self):
        """Test PartTypes CSV import - upsert existing record."""
        from Tracker.models import PartTypes

        # Create existing record
        existing = PartTypes.objects.create(
            name="Original Name",
            ERP_id="PT-002",
        )

        serializer = PartTypesCSVImportSerializer(
            user=self.user,
            mode=ImportMode.UPSERT,
        )

        row_data = {
            "name": "Updated Name",
            "ERP_id": "PT-002",  # Match by ERP_id
        }

        instance, created, warnings = serializer.import_row(row_data)

        self.assertFalse(created)  # Should be update
        self.assertEqual(instance.id, existing.id)
        self.assertEqual(instance.name, "Updated Name")

    def test_serializer_required_fields_validation(self):
        """Test required fields validation."""
        serializer = PartTypesCSVImportSerializer(
            user=self.user,
            mode=ImportMode.CREATE,
        )

        row_data = {
            "ERP_id": "PT-003",
            # Missing required 'name' field
        }

        from rest_framework.serializers import ValidationError
        with self.assertRaises(ValidationError):
            serializer.import_row(row_data)


class ImportValueReadingTests(TestCase):
    """Values an import can't read, or can read two ways, are reported on the row.

    Both used to be silent: an unreadable value was dropped (the field kept its old
    value and the row reported success), and `03/04/2026` became 4 March for a file
    whose author meant 3 April."""

    def _transform(self, row):
        s = OrdersCSVImportSerializer(data={}, mode="create")
        s.warnings = []
        return s.transform_data(row), s.warnings

    def test_an_ambiguous_date_is_read_month_first_and_says_so(self):
        out, warnings = self._transform({"estimated_completion": "03/04/2026"})
        self.assertEqual(str(out["estimated_completion"]), "2026-03-04")
        self.assertTrue(any("month first" in w for w in warnings), warnings)

    def test_an_unambiguous_date_needs_no_warning(self):
        _, warnings = self._transform({"estimated_completion": "2026-04-03"})
        self.assertEqual(warnings, [])

    def test_an_unreadable_value_is_reported_not_dropped_silently(self):
        out, warnings = self._transform({"estimated_completion": "next tuesday"})
        self.assertNotIn("estimated_completion", out)
        self.assertTrue(any("couldn't read" in w for w in warnings), warnings)


class CsvImportApiTests(APITestCase):
    """The import/export endpoints over HTTP.

    These used to call `/api/part-types/...` -- a route that doesn't exist -- with every
    assertion behind `if response.status_code == 200`, so each passed on a 404 and the
    suite never exercised the endpoints at all. That is how a crash in every export
    (2026-09-21 to -29) went unnoticed. They now hit the real routes and assert."""

    @classmethod
    def setUpTestData(cls):
        from Tracker.models import Tenant
        from Tracker.utils.tenant_context import reset_current_tenant, set_current_tenant_id
        cls.tenant = Tenant.objects.create(name="CSV Shop", slug="csv-shop")
        token = set_current_tenant_id(cls.tenant.id)
        try:
            cls.user = User.objects.create_user(
                username="csvadmin", email="csvadmin@example.com", password="testpass123",
                tenant=cls.tenant, is_staff=True)
            cls.user.is_superuser = True
            cls.user.save(update_fields=["is_superuser"])
        finally:
            reset_current_tenant(token)

    def setUp(self):
        from Tracker.utils.tenant_context import set_current_tenant_id
        self.client.force_authenticate(user=self.user)
        self.client.credentials(HTTP_X_TENANT_ID=str(self.tenant.id))
        self._token = set_current_tenant_id(self.tenant.id)

    def tearDown(self):
        from Tracker.utils.tenant_context import reset_current_tenant
        reset_current_tenant(self._token)

    def _upload(self, content, mode="create", user=None):
        if user is not None:
            self.client.force_authenticate(user=user)
        f = io.BytesIO(content.encode("utf-8"))
        f.name = "import.csv"
        return self.client.post("/api/PartTypes/import/", {"file": f, "mode": mode},
                                format="multipart")

    def test_import_template_endpoint_csv(self):
        response = self.client.get("/api/PartTypes/import-template/csv/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertIn("attachment", response.get("Content-Disposition", ""))

    def test_import_template_endpoint_xlsx(self):
        response = self.client.get("/api/PartTypes/import-template/xlsx/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheet", response["Content-Type"])

    def test_export_endpoint_csv(self):
        response = self.client.get("/api/PartTypes/export/csv/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], "text/csv")

    def test_export_endpoint_xlsx(self):
        response = self.client.get("/api/PartTypes/export/xlsx/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheet", response["Content-Type"])

    def test_import_endpoint_no_file(self):
        response = self.client.post("/api/PartTypes/import/")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_import_endpoint_with_csv(self):
        from Tracker.models import PartTypes
        response = self._upload("name,ERP_id\nTest Part,PT-TEST")
        self.assertEqual(response.status_code, status.HTTP_207_MULTI_STATUS, response.content)
        self.assertEqual(response.json()["summary"]["created"], 1, response.content)
        self.assertTrue(PartTypes.objects.filter(name="Test Part").exists())

    # --- what an import may write ---------------------------------------------------

    def test_an_import_cannot_move_a_record_to_another_tenant(self):
        """Columns were setattr'd onto the row, so `tenant_id` on an update moved it."""
        from Tracker.models import PartTypes, Tenant
        other = Tenant.objects.create(name="Elsewhere", slug="elsewhere")
        pt = PartTypes.objects.create(tenant=self.tenant, name="Stays Home")
        response = self._upload(f"id,name,tenant_id\n{pt.id},Stays Home,{other.id}", mode="update")
        self.assertEqual(response.status_code, status.HTTP_207_MULTI_STATUS, response.content)
        self.assertEqual(response.json()["summary"]["errors"], 1)
        pt.refresh_from_db()
        self.assertEqual(pt.tenant_id, self.tenant.id)

    def test_an_import_cannot_write_a_field_the_api_does_not_accept(self):
        from Tracker.models import PartTypes
        pt = PartTypes.objects.create(tenant=self.tenant, name="Versioned")
        response = self._upload(f"id,name,is_current_version\n{pt.id},Versioned,false", mode="update")
        self.assertEqual(response.json()["summary"]["errors"], 1, response.content)
        pt.refresh_from_db()
        self.assertTrue(pt.is_current_version)

    def test_a_read_only_column_is_ignored_with_a_warning(self):
        """An export carries columns the API shows but won't take. Ignored, not refused,
        so an exported file imports back — and never written."""
        from Tracker.models import PartTypes
        pt = PartTypes.objects.create(tenant=self.tenant, name="Shown")
        r = self._upload(f"id,name,ID_prefix,is_current_version\n{pt.id},Shown,SH,true", mode="update")
        body = r.json()
        self.assertEqual(body["summary"]["updated"], 1, r.content)
        self.assertIn("is_current_version", str(body["results"][0].get("warnings")), r.content)
        # A content edit versions a part type, as the API's does: read the current one.
        current = PartTypes.objects.get(name="Shown", is_current_version=True)
        self.assertEqual(current.ID_prefix, "SH")

    def test_a_column_that_names_no_field_is_still_refused(self):
        r = self._upload("name,colour\nTypo,red", mode="create")
        self.assertIn("colour", str(r.json()["results"][0]["errors"]), r.content)

    def test_update_modes_need_change_permission(self):
        """The POST gate checks ADD only; update/upsert rewrite existing rows."""
        from django.contrib.auth.models import Permission
        from Tracker.models import PartTypes, TenantGroup, UserRole
        adder = User.objects.create_user(username="adder", email="adder@example.com",
                                         password="x", tenant=self.tenant)
        group = TenantGroup.objects.create(tenant=self.tenant, name="Adders", is_custom=True)
        group.permissions.add(*Permission.objects.filter(
            codename__in=["add_parttypes", "view_parttypes"]))
        UserRole.objects.create(user=adder, group=group)
        pt = PartTypes.objects.create(tenant=self.tenant, name="Protected")
        r = self._upload(f"id,name\n{pt.id},Overwritten", mode="upsert", user=adder)
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN, r.content)
        pt.refresh_from_db()
        self.assertEqual(pt.name, "Protected")
        r = self._upload("name\nBrand New", mode="create", user=adder)
        self.assertEqual(r.status_code, status.HTTP_207_MULTI_STATUS, r.content)

    def test_one_bad_row_does_not_take_the_rest_with_it(self):
        """All rows ran in one transaction with no savepoints: after a row failed, every
        later row could fail with it."""
        from Tracker.models import PartTypes
        existing = PartTypes.objects.create(tenant=self.tenant, name="Taken", ERP_id="TAKEN")
        response = self._upload(f"id,name\n{existing.id},Taken\n,Fresh One", mode="create")
        summary = response.json()["summary"]
        self.assertEqual((summary["errors"], summary["created"]), (1, 1), response.content)
        self.assertTrue(PartTypes.objects.filter(name="Fresh One").exists())

    def test_import_status_is_not_readable_by_task_id_alone(self):
        response = self.client.get(f"/api/PartTypes/import-status/{uuid4()}/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    # --- formula injection ----------------------------------------------------------

    def test_exports_never_carry_a_live_formula(self):
        from openpyxl import load_workbook
        from Tracker.models import PartTypes
        PartTypes.objects.create(tenant=self.tenant, name='=HYPERLINK("http://x","go")')
        body = self.client.get("/api/PartTypes/export/csv/?fields=name").content.decode("utf-8-sig")
        self.assertIn(chr(39) + "=HYPERLINK", body)
        wb = load_workbook(io.BytesIO(self.client.get("/api/PartTypes/export/xlsx/?fields=name").content))
        cells = [c for row in wb["Data"].iter_rows(min_row=2) for c in row if c.value]
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == "s" for c in cells),
                        [(c.value, c.data_type) for c in cells])

    # --- round trips: the files we hand out must import back -------------------------

    def _post_file(self, url, content, name, mode):
        f = io.BytesIO(content)
        f.name = name
        return self.client.post(url, {"file": f, "mode": mode}, format="multipart")

    def test_a_filled_in_xlsx_template_imports_only_what_was_filled_in(self):
        """Required columns are headed `name*`, and the template put a hint row and a
        sample row under the header — the star matched nothing, and the two rows were
        imported as data."""
        from openpyxl import load_workbook
        from Tracker.models import PartTypes
        wb = load_workbook(io.BytesIO(self.client.get("/api/PartTypes/import-template/xlsx/").content))
        ws = wb["Data"]
        headers = [c.value for c in ws[1]]
        self.assertIn("name*", headers)
        ws.cell(row=2, column=headers.index("name*") + 1, value="From Template")
        out = io.BytesIO(); wb.save(out)
        r = self._post_file("/api/PartTypes/import/", out.getvalue(), "filled.xlsx", "create")
        self.assertEqual(r.json()["summary"], {"total": 1, "created": 1, "updated": 0, "errors": 0}, r.content)
        self.assertTrue(PartTypes.objects.filter(name="From Template").exists())
        self.assertFalse(PartTypes.objects.filter(name__in=["Part type name", "Widget Assembly"]).exists())

    def test_a_filled_in_csv_template_imports(self):
        """Our CSV template (like Excel's "CSV UTF-8") starts with a byte-order mark,
        which hid the first column."""
        from Tracker.models import PartTypes
        template = self.client.get("/api/PartTypes/import-template/csv/").content
        self.assertTrue(template.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(len(template.decode("utf-8-sig").strip().splitlines()), 1)
        header = template.decode("utf-8-sig").strip().split(",")
        row = ",".join("From CSV" if h == "name*" else "" for h in header)
        r = self._post_file("/api/PartTypes/import/", template + (row + "\n").encode(), "t.csv", "create")
        self.assertEqual(r.json()["summary"]["created"], 1, r.content)
        self.assertTrue(PartTypes.objects.filter(name="From CSV").exists())

    # --- the dialog's column mapping ------------------------------------------------

    def _upload_mapped(self, content, mapping, mode="create"):
        import json
        f = io.BytesIO(content.encode("utf-8"))
        f.name = "import.csv"
        return self.client.post("/api/PartTypes/import/", {
            "file": f, "mode": mode, "column_mapping": json.dumps(mapping)}, format="multipart")

    def test_a_manual_remap_of_a_spelled_out_header_is_applied(self):
        """The dialog keys its choices by the header as written; the parser looked them
        up by the normalized header only, so every manual remap was ignored."""
        from Tracker.models import PartTypes
        r = self._upload_mapped("Product Title\nRemapped", {"Product Title": "name"})
        self.assertEqual(r.json()["summary"]["created"], 1, r.content)
        self.assertTrue(PartTypes.objects.filter(name="Remapped").exists())

    def test_a_skipped_column_is_not_read(self):
        from Tracker.models import PartTypes
        r = self._upload_mapped("name,Internal Notes\nSkipper,ignore me",
                                {"name": "name", "Internal Notes": "_skip_"})
        self.assertEqual(r.json()["summary"]["created"], 1, r.content)
        self.assertTrue(PartTypes.objects.filter(name="Skipper").exists())

    def test_an_id_column_updates_that_record(self):
        """An exported file's `id` column is how a re-import finds its records; it was
        refused as a field the API doesn't accept."""
        from Tracker.models import PartTypes
        pt = PartTypes.objects.create(tenant=self.tenant, name="By Id")
        r = self._upload(f"id,ID_prefix\n{pt.id},BI", mode="update")
        self.assertEqual(r.json()["summary"]["updated"], 1, r.content)
        # A content edit versions a part type, as the API's does: read the current one.
        current = PartTypes.objects.get(name="By Id", is_current_version=True)
        self.assertEqual(current.ID_prefix, "BI")

    def test_the_preview_offers_id_to_match_on(self):
        f = io.BytesIO(b"id,name\n,x")
        f.name = "p.csv"
        r = self.client.post("/api/PartTypes/import-preview/", {"file": f}, format="multipart")
        body = r.json()
        self.assertIn("id", [m["name"] for m in body["model_fields"]], r.content)
        self.assertEqual(body["columns"][0]["mapped_to"], "id")

    # --- which existing record an import matches -----------------------------------

    def test_a_voided_record_is_not_matched(self):
        """`.objects` does not exclude archived rows, so an import used to upsert onto a
        voided record — or refuse a create because a voided one 'already exists'."""
        from Tracker.models import PartTypes
        PartTypes.objects.create(tenant=self.tenant, name="Retired", archived=True)
        response = self._upload("name\nRetired", mode="create")
        self.assertEqual(response.json()["summary"]["created"], 1, response.content)

    def test_a_name_matching_two_records_is_refused_not_guessed(self):
        from Tracker.models import PartTypes
        PartTypes.objects.create(tenant=self.tenant, name="Twin", ERP_id="TWIN-A")
        PartTypes.objects.create(tenant=self.tenant, name="Twin", ERP_id="TWIN-B")
        response = self._upload("name,ID_prefix\nTwin,TW", mode="upsert")
        body = response.json()
        self.assertEqual(body["summary"]["errors"], 1, response.content)
        self.assertIn("more than one", str(body["results"][0]["errors"]))
        self.assertFalse(PartTypes.objects.filter(ID_prefix="TW").exists())

    # --- templates and previews offer only what an import accepts -------------------

    def test_the_template_offers_only_importable_columns(self):
        response = self.client.get("/api/PartTypes/import-template/csv/")
        header = response.content.decode("utf-8-sig").splitlines()[0].lower()
        self.assertIn("name", header)
        for never in ("tenant", "is_current_version", "created_at"):
            self.assertNotIn(never, header, header)

    # --- limits ---------------------------------------------------------------------

    def test_an_old_format_xls_says_what_to_do(self):
        f = io.BytesIO(b"not really a workbook")
        f.name = "legacy.xls"
        r = self.client.post("/api/PartTypes/import/", {"file": f, "mode": "create"},
                             format="multipart")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.content)
        self.assertIn(".xlsx", r.json()["detail"])

    def test_too_many_rows_is_refused_before_any_is_written(self):
        from Tracker.models import PartTypes
        from Tracker.services.csv_utils import MAX_UPLOAD_ROWS
        rows = "\n".join(f"Bulk {i}" for i in range(MAX_UPLOAD_ROWS + 1))
        r = self._upload("name\n" + rows, mode="create")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.content[:200])
        self.assertFalse(PartTypes.objects.filter(name__startswith="Bulk ").exists())


class PartsStepImportTests(APITestCase):
    """A parts import may place parts at a step — loading in-flight work — resolved
    within the part's own process, since step names repeat across processes."""

    @classmethod
    def setUpTestData(cls):
        from Tracker.models import (
            PartTypes, Processes, ProcessStep, Steps, Tenant, WorkOrder, WorkOrderStatus,
        )
        from Tracker.utils.tenant_context import reset_current_tenant, set_current_tenant_id
        cls.tenant = Tenant.objects.create(name="Step Shop", slug="step-shop")
        token = set_current_tenant_id(cls.tenant.id)
        try:
            cls.user = User.objects.create_user(
                username="stepadmin", email="stepadmin@example.com", password="x",
                tenant=cls.tenant, is_staff=True)
            cls.user.is_superuser = True
            cls.user.save(update_fields=["is_superuser"])
            cls.pt = PartTypes.objects.create(tenant=cls.tenant, name="Pump")
            procs = {}
            for pname in ("Build", "Repair"):
                proc = Processes.objects.create(tenant=cls.tenant, name=pname,
                                                part_type=cls.pt, status="APPROVED")
                steps = {}
                for i, sname in enumerate(("Cut", "Assembly", "Test"), start=1):
                    st = Steps.objects.create(tenant=cls.tenant, part_type=cls.pt, name=sname)
                    ProcessStep.objects.create(process=proc, step=st, order=i)
                    steps[sname] = st
                procs[pname] = (proc, steps)
            cls.build, cls.build_steps = procs["Build"]
            cls.repair_steps = procs["Repair"][1]
            cls.wo = WorkOrder.objects.create(
                tenant=cls.tenant, ERP_id="WO-STEP", quantity=1,
                workorder_status=WorkOrderStatus.IN_PROGRESS, process=cls.build)
        finally:
            reset_current_tenant(token)

    def setUp(self):
        from Tracker.utils.tenant_context import set_current_tenant_id
        self.client.force_authenticate(user=self.user)
        self.client.credentials(HTTP_X_TENANT_ID=str(self.tenant.id))
        self._token = set_current_tenant_id(self.tenant.id)

    def tearDown(self):
        from Tracker.utils.tenant_context import reset_current_tenant
        reset_current_tenant(self._token)

    def _upload(self, content, mode="upsert"):
        f = io.BytesIO(content.encode("utf-8"))
        f.name = "parts.csv"
        return self.client.post("/api/Parts/import/", {"file": f, "mode": mode},
                                format="multipart")

    def test_a_new_part_is_placed_at_its_processs_step_by_name(self):
        """'Assembly' exists in both processes; the work order's process decides."""
        from Tracker.models import Parts
        r = self._upload("ERP_id,part_type,work_order,step\nP-1,Pump,WO-STEP,Assembly", mode="create")
        self.assertEqual(r.json()["summary"]["created"], 1, r.content)
        self.assertEqual(Parts.objects.get(ERP_id="P-1").step_id, self.build_steps["Assembly"].id)

    def test_a_step_outside_the_process_is_refused(self):
        r = self._upload("ERP_id,part_type,work_order,step\nP-2,Pump,WO-STEP,Paint", mode="create")
        body = r.json()
        self.assertEqual(body["summary"]["errors"], 1, r.content)
        self.assertIn("not a step of", str(body["results"][0]["errors"]))

    def test_an_existing_part_can_be_moved_to_a_step(self):
        from Tracker.models import Parts
        part = Parts.objects.create(tenant=self.tenant, ERP_id="P-3", part_type=self.pt,
                                    work_order=self.wo, step=self.build_steps["Cut"])
        r = self._upload(f"id,step\n{part.id},Test", mode="update")
        self.assertEqual(r.json()["summary"]["updated"], 1, r.content)
        part.refresh_from_db()
        self.assertEqual(part.step_id, self.build_steps["Test"].id)

    def test_a_core_part_is_not_moved_by_import(self):
        from datetime import date
        from Tracker.services.reman.core_part import create_core
        core = create_core(tenant=self.tenant, core_number="CORE-STEP", core_type=self.pt,
                           received_date=date.today(), received_by=self.user,
                           condition_grade="B", work_order=self.wo,
                           step=self.build_steps["Cut"])
        r = self._upload(f"id,step\n{core.part_id},Test", mode="update")
        self.assertEqual(r.json()["summary"]["errors"], 1, r.content)
        core.part.refresh_from_db()
        self.assertEqual(core.part.step_id, self.build_steps["Cut"].id)

    def test_a_create_still_needs_its_required_columns(self):
        """Required columns are relaxed for updates only."""
        r = self._upload("ERP_id,work_order\nP-4,WO-STEP", mode="create")
        self.assertIn("part_type", str(r.json()["results"][0]["errors"]), r.content)

    def test_the_parts_template_offers_step(self):
        header = self.client.get("/api/Parts/import-template/csv/").content.decode("utf-8-sig")
        self.assertIn("step", header.splitlines()[0].lower())

    def test_an_exported_file_imports_back_as_an_update(self):
        """Export headers are labels — `Part Type Name`, `Part Type (auto)`, `Created
        At` — and none of them mapped back, so every row of a re-import failed."""
        from Tracker.models import Parts
        part = Parts.objects.create(tenant=self.tenant, ERP_id="P-RT", part_type=self.pt,
                                    work_order=self.wo, step=self.build_steps["Cut"])
        exported = self.client.get("/api/Parts/export/xlsx/").content
        # Uploaded as it is, straight to the API: read-only columns are ignored.
        f = io.BytesIO(exported); f.name = "parts-export.xlsx"
        raw = self.client.post("/api/Parts/import/", {"file": f, "mode": "update"}, format="multipart")
        self.assertEqual(raw.json()["summary"]["updated"], 1, raw.content)
        # And as the dialog does it: preview, keep its mapping (unmatched columns skipped).
        import json
        f = io.BytesIO(exported); f.name = "parts-export.xlsx"
        cols = self.client.post("/api/Parts/import-preview/", {"file": f},
                                format="multipart").json()["columns"]
        mapping = {c["original"]: c["mapped_to"] or "_skip_" for c in cols}
        self.assertEqual(mapping.get("part_type_name"), "part_type", mapping)
        self.assertEqual(mapping.get("step_ref"), "step", mapping)
        f = io.BytesIO(exported); f.name = "parts-export.xlsx"
        r = self.client.post("/api/Parts/import/", {"file": f, "mode": "update",
                                                    "column_mapping": json.dumps(mapping)},
                             format="multipart")
        body = r.json()
        self.assertEqual(body["summary"]["errors"], 0, r.content)
        self.assertEqual(body["summary"]["updated"], 1, r.content)
        part.refresh_from_db()
        self.assertEqual(part.part_type_id, self.pt.id)

    def test_the_preview_maps_an_exported_fk_name_column_to_the_fk(self):
        f = io.BytesIO(b"Part Type Name,Part Type (auto)\nPump,\n"); f.name = "p.csv"
        cols = self.client.post("/api/Parts/import-preview/", {"file": f},
                                format="multipart").json()["columns"]
        self.assertEqual([c["mapped_to"] for c in cols], ["part_type", None])

    def test_the_parts_template_explains_step(self):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.client.get("/api/Parts/import-template/xlsx/").content))
        text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
        self.assertIn("work-order process", text)


class AutomaticIntrospectionTests(TestCase):
    """Tests for automatic model introspection."""

    def test_template_from_model(self):
        """Test automatic template generation from model."""
        from Tracker.models import PartTypes

        template = template_from_model(PartTypes)

        self.assertEqual(template.model_name, "PartTypes")
        # Should have auto-detected fields
        field_names = [f.name for f in template.fields]
        self.assertIn("name", field_names)

    def test_introspect_model_finds_fk_fields(self):
        """Test that introspection finds foreign key relationships."""
        from Tracker.models import Parts

        fields = introspect_model(Parts)
        field_names = [f.name for f in fields]

        # Should find FK fields
        self.assertIn("part_type", field_names)
        self.assertIn("order", field_names)

        # FK fields should have fk_model set
        part_type_field = next(f for f in fields if f.name == "part_type")
        self.assertIsNotNone(part_type_field.fk_model)

    def test_create_import_serializer_for_model(self):
        """Test dynamic serializer creation from model."""
        from Tracker.models import PartTypes

        serializer_class = create_import_serializer_for_model(PartTypes)

        # Should be a subclass of BaseCSVImportSerializer
        self.assertTrue(issubclass(serializer_class, BaseCSVImportSerializer))

        # Should have Meta with model
        self.assertEqual(serializer_class.Meta.model, PartTypes)

        # Should have lookup_fields
        self.assertIn('id', serializer_class.Meta.lookup_fields)

    def test_get_or_create_import_serializer_uses_registry(self):
        """Test that registered serializers are returned first."""
        from Tracker.models import PartTypes

        serializer_class = get_or_create_import_serializer(PartTypes)

        # Should return the registered serializer, not a new one
        self.assertEqual(serializer_class, PartTypesCSVImportSerializer)

    def test_get_or_create_import_serializer_creates_new(self):
        """Test that new serializers are created for unregistered models."""
        from Tracker.models import EquipmentType

        serializer_class = get_or_create_import_serializer(EquipmentType)

        # Should be a subclass of BaseCSVImportSerializer
        self.assertTrue(issubclass(serializer_class, BaseCSVImportSerializer))


class ForeignKeyResolutionTests(TenantContextMixin, TestCase):
    """Tests for FK resolution in import serializers."""

    @classmethod
    def setUpTestData(cls):
        """Create test data with related models."""
        from Tracker.models import PartTypes, Companies, Tenant

        cls.tenant = Tenant.objects.create(name="FK Test Tenant", slug="fk-test")
        cls.set_tenant_context_class(cls.tenant)

        cls.user = User.objects.create_user(
            username="testuser",
            email="test@example.com",
            password="testpass123",
            tenant=cls.tenant,
        )

        cls.part_type = PartTypes.objects.create(
            name="Widget",
            ERP_id="PT-WIDGET",
        )

        cls.company = Companies.objects.create(
            name="Acme Corp",
        )

    def test_resolve_fk_by_name(self):
        """Test FK resolution by name."""
        from Tracker.models import PartTypes

        serializer = PartsCSVImportSerializer(user=self.user)
        resolved = serializer.resolve_fk(
            "part_type",
            "Widget",
            PartTypes,
            ["name", "ERP_id", "id"],
        )

        self.assertEqual(resolved, self.part_type)

    def test_resolve_fk_by_erp_id(self):
        """Test FK resolution by ERP_id."""
        from Tracker.models import PartTypes

        serializer = PartsCSVImportSerializer(user=self.user)
        resolved = serializer.resolve_fk(
            "part_type",
            "PT-WIDGET",
            PartTypes,
            ["name", "ERP_id", "id"],
        )

        self.assertEqual(resolved, self.part_type)

    def test_resolve_fk_by_uuid(self):
        """Test FK resolution by UUID."""
        from Tracker.models import PartTypes

        serializer = PartsCSVImportSerializer(user=self.user)
        resolved = serializer.resolve_fk(
            "part_type",
            str(self.part_type.id),
            PartTypes,
            ["name", "ERP_id", "id"],
        )

        self.assertEqual(resolved, self.part_type)

    def test_resolve_fk_case_insensitive(self):
        """Test FK resolution is case-insensitive."""
        from Tracker.models import PartTypes

        serializer = PartsCSVImportSerializer(user=self.user)
        resolved = serializer.resolve_fk(
            "part_type",
            "WIDGET",  # Different case
            PartTypes,
            ["name", "ERP_id", "id"],
        )

        self.assertEqual(resolved, self.part_type)

    def test_resolve_fk_not_found(self):
        """Test FK resolution returns None when not found."""
        from Tracker.models import PartTypes

        serializer = PartsCSVImportSerializer(user=self.user)
        resolved = serializer.resolve_fk(
            "part_type",
            "Nonexistent Part Type",
            PartTypes,
            ["name", "ERP_id", "id"],
        )

        self.assertIsNone(resolved)
