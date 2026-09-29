"""The data export (`DataExportMixin`, `/api/<Model>/export/<csv|xlsx>/`).

An export reads the MODEL, not the serializer, so it is a second way out of the API for
every field a serializer hides. These pin that it shows no more than the API does, that
`?fields=` cannot reach past that, and that related columns land on their own rows.
Found 2026-09-29: the User export included password hashes, and `?fields=` could pull
`<fk>__password` through any model with a user relation.
"""
import csv
import io
from datetime import date

from openpyxl import load_workbook
from rest_framework.test import APIClient
from django.test import TestCase

from Tracker.models import PartTypes, Tenant, User
from Tracker.services.reman.core_part import create_core
from Tracker.utils.tenant_context import reset_current_tenant, set_current_tenant_id


class DataExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(name="Export Shop", slug="export-shop")
        token = set_current_tenant_id(cls.tenant.id)
        try:
            cls.user = User.objects.create_user(
                username="exporter", email="exporter@test.com", password="s3cret-pass",
                tenant=cls.tenant, is_staff=True)
            cls.user.is_superuser = True
            cls.user.save(update_fields=['is_superuser'])
            cls.injector = PartTypes.objects.create(tenant=cls.tenant, name="Injector")
            cls.turbo = PartTypes.objects.create(tenant=cls.tenant, name="Turbo")
            # Same received date on every core: an ordering with ties, which is where
            # position-matched related columns used to go wrong.
            for i in range(6):
                create_core(tenant=cls.tenant, core_number=f"EXP-{i}",
                            core_type=cls.injector if i % 2 else cls.turbo,
                            received_date=date(2026, 9, 1), received_by=cls.user,
                            condition_grade='B')
        finally:
            reset_current_tenant(token)

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.client.credentials(HTTP_X_TENANT_ID=str(self.tenant.id))
        self._token = set_current_tenant_id(self.tenant.id)

    def tearDown(self):
        reset_current_tenant(self._token)

    def _csv(self, url):
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return list(csv.DictReader(io.StringIO(r.content.decode('utf-8-sig'))))

    def test_the_user_export_carries_no_password_hash(self):
        rows = self._csv("/api/User/export/csv/")
        self.assertTrue(rows)
        self.assertNotIn('password', {k.lower() for k in rows[0]})
        self.assertFalse(any(self.user.password in v for row in rows for v in row.values()))

    def test_fields_cannot_reach_a_password_through_a_relation(self):
        r = self.client.get("/api/Cores/export/csv/?fields=core_number,received_by__password")
        self.assertEqual(r.status_code, 400, r.content[:300])
        self.assertIn('received_by__password', str(r.json()))

    def test_fields_cannot_ask_for_what_the_serializer_does_not_show(self):
        # `received_by__email` is a real path, but the core serializer does not show it.
        r = self.client.get("/api/Cores/export/csv/?fields=core_number,received_by__email")
        self.assertEqual(r.status_code, 400, r.content[:300])

    def test_related_columns_land_on_their_own_rows(self):
        rows = self._csv("/api/Cores/export/csv/?fields=core_number,core_type__name")
        by_number = {r['Core Number'] if 'Core Number' in r else r['core_number']: r for r in rows}
        for i in range(6):
            row = by_number[f"EXP-{i}"]
            name = row.get('Core Type Name') or row.get('core_type__name')
            self.assertEqual(name, "Injector" if i % 2 else "Turbo", f"EXP-{i}")

    def test_xlsx_builds_and_only_references_exported_relations(self):
        r = self.client.get("/api/Cores/export/xlsx/?fields=core_number,core_type,core_type__name")
        self.assertEqual(r.status_code, 200, r.content[:300])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn('Data', wb.sheetnames)
        # A reference sheet for the exported core type — and none for the core's other
        # relations (the receiving user, the customer, ...), which were not exported.
        self.assertFalse([n for n in wb.sheetnames if 'received' in n.lower() or 'user' in n.lower()],
                         wb.sheetnames)


class AccessLogPayloadTests(TestCase):
    """`record_access` puts its payload on the compliance log record. A payload key that
    a LogRecord reserves (`filename`, `module`, ...) made logging raise — which is how
    every export 500'd from 2026-09-21 until this."""

    def test_a_reserved_key_in_the_payload_does_not_crash_the_request(self):
        from Tracker.services.core.access_log import _log_safe
        safe = _log_safe({'filename': 'x.csv', 'module': 'm', 'row_count': 3})
        self.assertEqual(safe, {'payload_filename': 'x.csv', 'payload_module': 'm',
                                'row_count': 3})
