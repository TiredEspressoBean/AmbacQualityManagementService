"""Fixes from the receiving review (2026-10-01): lot integrity, splits and stuck lots,
the numbers planning and the dock read, and the expected-receipts sheet."""
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase

from Tracker.utils.tenant_context import reset_current_tenant, set_current_tenant_id

User = get_user_model()


class _Fixture(APITestCase):
    @classmethod
    def setUpTestData(cls):
        from Tracker.models import Companies, Material, Tenant
        cls.tenant = Tenant.objects.create(name="Review", slug="receiving-review-1001")
        token = set_current_tenant_id(cls.tenant.id)
        try:
            cls.user = User.objects.create_user(username="rv", email="rv@x.test", password="x",
                                                tenant=cls.tenant, is_staff=True)
            cls.user.is_superuser = True
            cls.user.save(update_fields=["is_superuser"])
            cls.acme = Companies.objects.create(tenant=cls.tenant, name="Acme")
            cls.fleet = Companies.objects.create(tenant=cls.tenant, name="Fleet", is_customer=True)
            cls.seal = Material.objects.create(tenant=cls.tenant, name="Seal", part_number="SEAL",
                                               unit_of_measure="EA")
            cls.bar = Material.objects.create(tenant=cls.tenant, name="Bar", unit_of_measure="KG")
        finally:
            reset_current_tenant(token)

    def setUp(self):
        self.client.force_authenticate(user=self.user)
        self.client.credentials(HTTP_X_TENANT_ID=str(self.tenant.id))
        self._token = set_current_tenant_id(self.tenant.id)

    def tearDown(self):
        reset_current_tenant(self._token)

    def _lot(self, number="L-1", status="ACCEPTED", qty="100", material=None, **extra):
        from Tracker.models import MaterialLot
        return MaterialLot.objects.create(
            tenant=self.tenant, lot_number=number, material=material or self.seal,
            supplier=self.acme, quantity=Decimal(qty), quantity_remaining=Decimal(qty),
            unit_of_measure=(material or self.seal).unit_of_measure, status=status,
            received_date=date.today(), **extra)

    def _expected(self, po="4500", line="10", qty="50"):
        from Tracker.services.mes.material_lot import record_expected_receipt
        return record_expected_receipt(tenant=self.tenant, quantity=Decimal(qty), material=self.seal,
                                       promised_date=date.today() + timedelta(days=3),
                                       erp_po_number=po, erp_po_line=line)


class LotIntegrityTests(_Fixture):
    def test_status_cant_be_written_through_the_api(self):
        from Tracker.models import MaterialLot
        resp = self.client.post("/api/MaterialLots/", {
            "lot_number": "W-1", "material": str(self.seal.id), "quantity": "5",
            "unit_of_measure": "EA", "received_date": str(date.today()), "status": "ACCEPTED",
        }, format="json")
        self.assertEqual(resp.status_code, 201, resp.content)
        # Routed as any receipt is (dock-to-stock here: no receiving plan) — not because
        # the request said so. A RETURNED lot can't be patched back into stock.
        returned = self._lot("R-1", status="RETURNED")
        self.client.patch(f"/api/MaterialLots/{returned.id}/", {"status": "ACCEPTED"}, format="json")
        self.assertEqual(MaterialLot.objects.get(pk=returned.id).status, "RETURNED")

    def test_quantity_changes_through_adjust_not_an_edit(self):
        lot = self._lot()
        resp = self.client.patch(f"/api/MaterialLots/{lot.id}/", {"quantity": "80"}, format="json")
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Adjust quantity", str(resp.json()))
        # The same quantity sent back (a form saving every field) is no change.
        ok = self.client.patch(f"/api/MaterialLots/{lot.id}/", {"quantity": "100", "storage_location": "B2"},
                               format="json")
        self.assertEqual(ok.status_code, 200, ok.content)

    def test_receipts_are_numbered_by_us_and_keep_the_suppliers_lot(self):
        first, second = self._expected(line="10"), self._expected(line="20")
        # The same supplier lot twice — a back-order of one lot — both land.
        for lot in (first, second):
            resp = self.client.post(f"/api/MaterialLots/{lot.id}/receive/",
                                    {"supplier_lot_number": "240915"}, format="json")
            self.assertEqual(resp.status_code, 200, resp.content)
        first.refresh_from_db(), second.refresh_from_db()
        year = date.today().year
        self.assertEqual((first.lot_number, second.lot_number),
                         (f"LOT-{year}-00001", f"LOT-{year}-00002"))
        self.assertEqual((first.supplier_lot_number, second.supplier_lot_number), ("240915", "240915"))

    def test_a_typed_lot_number_that_clashes_is_a_400_not_a_crash(self):
        self._lot("DUP-1")
        resp = self.client.post("/api/MaterialLots/", {
            "lot_number": "dup-1", "material": str(self.seal.id), "quantity": "5",
            "unit_of_measure": "EA", "received_date": str(date.today())}, format="json")
        self.assertEqual(resp.status_code, 400)
        self.assertIn("already in use", str(resp.json()))
        lot = self._expected()
        resp = self.client.post(f"/api/MaterialLots/{lot.id}/receive/", {"lot_number": "DUP-1"},
                                format="json")
        self.assertEqual(resp.status_code, 400)

    def test_batch_receive_numbers_blank_rows_and_refuses_a_repeated_one(self):
        from Tracker.models import MaterialLot
        row = {"material": str(self.seal.id), "quantity": "5", "unit_of_measure": "EA",
               "received_date": str(date.today())}
        resp = self.client.post("/api/MaterialLots/bulk_create/",
                                {"lots": [{**row, "lot_number": ""}, {**row}]}, format="json")
        self.assertEqual(resp.status_code, 201, resp.content)
        numbers = sorted(MaterialLot.objects.filter(id__in=resp.json()["created_lot_ids"])
                         .values_list("lot_number", flat=True))
        self.assertEqual(numbers, [f"LOT-{date.today().year}-0000{i}" for i in (1, 2)])
        resp = self.client.post("/api/MaterialLots/bulk_create/",
                                {"lots": [{**row, "lot_number": "X"}, {**row, "lot_number": "x"}]},
                                format="json")
        self.assertEqual(resp.status_code, 400)

    def test_counting_a_lot_to_zero_uses_it_up(self):
        from Tracker.services.mes.material_lot import adjust_quantity
        lot = adjust_quantity(self._lot(), new_quantity=Decimal(0), reason="Empty bin", user=self.user)
        self.assertEqual(lot.status, "CONSUMED")
        for status in ("REJECTED", "RETURNED"):
            with self.assertRaises(ValueError):
                adjust_quantity(self._lot(f"S-{status}", status=status), new_quantity=Decimal(5),
                                reason="x", user=self.user)


class SplitTests(_Fixture):
    def test_a_split_of_stock_is_stock_and_stays_the_owners(self):
        from Tracker.services.mes.material_lot import split_material_lot
        parent = self._lot(owner=self.fleet, promised_date=date.today())
        child = split_material_lot(parent, Decimal(30))
        self.assertEqual((child.status, child.owner_id, child.promised_date),
                         ("ACCEPTED", self.fleet.id, date.today()))

    def test_a_held_lot_splits_held_and_can_be_partly_rejected(self):
        from Tracker.services.mes.material_lot import split_material_lot
        held = self._lot("H-1", status="QUARANTINE", hold_reason="SUPPLIER_UNQUALIFIED")
        child = split_material_lot(held, Decimal(10))
        self.assertEqual((child.status, child.hold_reason), ("QUARANTINE", "SUPPLIER_UNQUALIFIED"))

    def test_a_received_lot_still_splits_to_received(self):
        from Tracker.services.mes.material_lot import split_material_lot
        child = split_material_lot(self._lot("RC-1", status="RECEIVED"), Decimal(10))
        self.assertEqual(child.status, "RECEIVED")


class DispositionTests(_Fixture):
    def _disposition(self, lot, kind):
        from Tracker.models import QuarantineDisposition
        return QuarantineDisposition.objects.create(
            tenant=self.tenant, material_lot=lot, disposition_type=kind, quantity=lot.quantity,
            current_state="IN_PROGRESS", severity="MINOR")

    def test_a_lot_cant_be_reworked(self):
        from Tracker.services.qms.disposition import decide_disposition
        lot = self._lot("RW-1", status="REJECTED")
        d = self._disposition(lot, "")
        for kind in ("REWORK", "REPAIR"):
            with self.assertRaises(ValueError):
                decide_disposition(d, disposition_type=kind, authorized_by=self.user,
                                   customer_approval={"reference": "C-1"})

    def test_completing_a_return_ships_it_back(self):
        from Tracker.models import MaterialLot
        from Tracker.services.qms.disposition import complete_disposition_resolution
        lot = self._lot("RTV-1", status="REJECTED")
        d = complete_disposition_resolution(self._disposition(lot, "RETURN_TO_SUPPLIER"), self.user)
        self.assertEqual(d.current_state, "CLOSED")
        lot = MaterialLot.objects.get(pk=lot.id)
        self.assertEqual((lot.status, lot.quantity_remaining), ("RETURNED", Decimal(0)))


class NumbersTests(_Fixture):
    def test_ppm_counts_our_pieces_only(self):
        from Tracker.models import QuarantineDisposition
        from Tracker.services.mes.dock_metrics import dock_metrics
        ours = self._lot("P-1", qty="1000")
        theirs = self._lot("P-2", qty="1000", owner=self.fleet)
        kilos = self._lot("P-3", qty="500", material=self.bar)
        for lot, qty in ((ours, 10), (theirs, 10), (kilos, 100)):
            QuarantineDisposition.objects.create(tenant=self.tenant, material_lot=lot,
                                                 disposition_type="SCRAP", quantity=qty)
        QuarantineDisposition.objects.create(
            tenant=self.tenant, material_lot=ours, disposition_type="RETURN_TO_SUPPLIER",
            quantity=1000, current_state="CLOSED", resolution_notes="Whole-lot reject declined: fine")
        m = dock_metrics(self.tenant)
        self.assertEqual(m["pieces_rejected"], 10.0)
        self.assertEqual(m["ppm_rejected"], 10_000)  # 10 of our 1,000 pieces

    def test_held_stock_isnt_incoming_supply(self):
        from Tracker.services.mes.requirements import _NOT_INCOMING_LOT_STATUSES
        self.assertIn("QUARANTINE", _NOT_INCOMING_LOT_STATUSES)

    def test_the_trace_follows_split_pieces(self):
        from Tracker.models import MaterialUsage, Parts, PartTypes, WorkOrder, WorkOrderStatus
        from Tracker.services.mes.lot_trace import trace_lot
        from Tracker.services.mes.material_lot import split_material_lot
        parent = self._lot("T-1")
        child = split_material_lot(parent, Decimal(40))
        wo = WorkOrder.objects.create(tenant=self.tenant, ERP_id="WO-1", quantity=1,
                                      workorder_status=WorkOrderStatus.IN_PROGRESS)
        part = Parts.objects.create(tenant=self.tenant, ERP_id="P-1", work_order=wo,
                                    part_type=PartTypes.objects.create(tenant=self.tenant, name="Pump"))
        MaterialUsage.objects.create(tenant=self.tenant, lot=child, part=part, work_order=wo,
                                     qty_consumed=Decimal(2), consumed_by=self.user)
        forward = trace_lot(parent)["forward"]
        self.assertEqual([(u["lot_number"], u["part"]["erp_id"]) for u in forward],
                         [(child.lot_number, "P-1")])


class ExpectedSheetTests(_Fixture):
    def test_a_day_month_date_is_flagged(self):
        from Tracker.services.mes.expected_receipt_import import import_expected_rows
        report = import_expected_rows(tenant=self.tenant, rows=[
            {"po": "77", "line": "1", "item": "SEAL", "quantity": "5", "promised": "03/04/2027"}])
        self.assertEqual(report["created"], 1, report)
        self.assertIn("month first", report["rows"][0]["detail"])


class ReceiptsExportTests(_Fixture):
    def test_a_delivery_counts_its_split_off_rejects_back(self):
        from Tracker.services.mes.material_lot import split_material_lot
        from Tracker.services.mes.receipt_export import receipt_rows
        d = self._lot("D-1", status="AWAITING_INSPECTION", qty="100",
                      erp_po_number="4500", erp_po_line="10", supplier_lot_number="S9")
        bad = split_material_lot(d, Decimal(15))
        bad.status = "REJECTED"
        bad.save(update_fields=["status"])
        d.status = "ACCEPTED"
        d.save(update_fields=["status"])
        self._lot("D-2", status="AWAITING_INSPECTION", qty="5", erp_po_number="4500", erp_po_line="20")
        self._lot("WALKIN", qty="3")
        rows = receipt_rows(self.tenant, date.today(), date.today(), with_po_only=True)
        by_line = {r["PO Line"]: r for r in rows}
        self.assertEqual(set(by_line), {"10", "20"})  # the walk-in left out, the split piece too
        first = by_line["10"]
        self.assertEqual((first["Received Qty"], first["Accepted"], first["Rejected"],
                          first["Awaiting Decision"], first["Decision"], first["Supplier Lot"]),
                         (Decimal(100), Decimal(85), Decimal(15), Decimal(0), "Partly rejected", "S9"))
        self.assertEqual(by_line["20"]["Decision"], "Awaiting decision")

    def test_the_sheet_downloads(self):
        from openpyxl import load_workbook
        import io
        self._lot("D-3", erp_po_number="4501", erp_po_line="1")
        today = str(date.today())
        resp = self.client.get(f"/api/MaterialLots/receipts-export/?start={today}&end={today}")
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(io.BytesIO(resp.content))["Receipts"]
        self.assertEqual(ws["A2"].value, "4501")
        self.assertEqual(self.client.get("/api/MaterialLots/receipts-export/?start=x").status_code, 400)


class MirrorTests(_Fixture):
    def test_a_received_line_the_sheet_shows_open_is_expected_again_and_flagged(self):
        from Tracker.models import MaterialLot
        from Tracker.services.mes.expected_receipt_import import import_expected_rows
        from Tracker.services.mes.material_lot import receive_expected_lot
        row = {"po": "88", "line": "1", "item": "SEAL", "quantity": "5", "promised": "2027-01-15"}
        import_expected_rows(tenant=self.tenant, rows=[row])
        receive_expected_lot(MaterialLot.objects.get(erp_po_number="88"), received_by=self.user)
        report = import_expected_rows(tenant=self.tenant, rows=[row])
        self.assertEqual(report["reopened"], 1, report)
        self.assertIn("If the ERP hasn't caught up", report["rows"][0]["detail"])
        self.assertEqual(MaterialLot.objects.filter(erp_po_number="88", status="ON_ORDER").count(), 1)


class LineageAndCancelTests(_Fixture):
    def test_a_split_lot_knows_the_delivery_its_paperwork_is_on(self):
        from Tracker.services.mes.material_lot import split_material_lot
        parent = self._lot("LIN-1")
        child = split_material_lot(parent, Decimal(10))
        grandchild = split_material_lot(child, Decimal(2))
        body = self.client.get(f"/api/MaterialLots/{grandchild.id}/").json()
        self.assertEqual([a["lot_number"] for a in body["lineage"]], [child.lot_number, "LIN-1"])

    def test_cancelling_an_expected_receipt(self):
        from Tracker.models import MaterialLot, RecordEdit
        lot = self._expected(po="4410", line="2")
        self.assertEqual(self.client.post(f"/api/MaterialLots/{lot.id}/cancel-expected/", {},
                                          format="json").status_code, 400)  # a reason is required
        resp = self.client.post(f"/api/MaterialLots/{lot.id}/cancel-expected/",
                                {"reason": "Already received"}, format="json")
        self.assertEqual(resp.status_code, 200, resp.content)
        lot = MaterialLot.objects.get(pk=lot.id)
        self.assertEqual((lot.status, lot.quantity_remaining), ("CANCELLED", Decimal(0)))
        self.assertTrue(RecordEdit.objects.filter(object_id=lot.id, reason="Already received").exists())
        # Not supply: requirements' incoming set leaves it out.
        from Tracker.services.mes.requirements import _NOT_INCOMING_LOT_STATUSES
        self.assertIn("CANCELLED", _NOT_INCOMING_LOT_STATUSES)

    def test_cancelling_a_backorder_remainder_closes_its_delivery(self):
        from Tracker.models import MaterialLot
        from Tracker.services.mes.material_lot import cancel_expected_receipt, receive_expected_lot
        first = self._expected(po="4411", line="1", qty="10")
        receive_expected_lot(first, received_by=self.user, quantity=Decimal(6), remainder="BACKORDERED")
        rest = MaterialLot.objects.get(erp_po_number="4411", status="ON_ORDER")
        cancel_expected_receipt(rest, user=self.user, reason="Supplier cancelled the balance")
        self.assertEqual(MaterialLot.objects.get(pk=first.id).short_receipt, "CLOSED")

    def test_a_cancelled_line_the_sheet_still_shows_is_expected_again_and_says_so(self):
        from Tracker.services.mes.expected_receipt_import import import_expected_rows
        from Tracker.services.mes.material_lot import cancel_expected_receipt
        row = {"po": "4412", "line": "1", "item": "SEAL", "quantity": "5", "promised": "2027-01-15"}
        import_expected_rows(tenant=self.tenant, rows=[row])
        from Tracker.models import MaterialLot
        cancel_expected_receipt(MaterialLot.objects.get(erp_po_number="4412"), user=self.user,
                                reason="Typed early")
        report = import_expected_rows(tenant=self.tenant, rows=[row])
        self.assertEqual(report["reopened"], 1, report)
        self.assertIn("was cancelled", report["rows"][0]["detail"])
