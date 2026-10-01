"""The master migration workbook — a new plant's data, one sheet per table, loaded at once.

Design: Documents/WORKBOOK_IMPORT_EXPORT_DESIGN.md. A conductor over the per-table
imports, not a second importer: each sheet runs through the import its own table
already has (the viewset's `CSVImportMixin` path, grouped imports included — a BOM's
lines, a ruleset's rules land together), or, for the few tables whose import is a
service, through that service.

- **In load order.** `SHEETS` is ordered so every sheet's rows name things an earlier
  sheet made. All of it runs in ONE transaction, so a part type created on the Part
  Types sheet is there for the BOMs sheet to name.
- **Dry run first.** The same run, rolled back at the end: every row's outcome and
  every error, before anything is kept.
- **All or nothing.** A real load keeps nothing unless every row of every sheet
  loaded. A half-loaded migration is worse than none — the next upload would be
  fixing a database rather than a spreadsheet.
- **Re-uploadable.** Every sheet upserts on its table's natural key, so the analyst
  fixes the workbook and uploads it again; versioned rows version, nothing is deleted.

Processes and steps are not on it: a routing is a flow, authored in the process
editor, before the workbook's work orders (which name their process) are loaded.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from django.db import transaction
from django.utils.module_loading import import_string

# The most rows one workbook may carry, all sheets together. Larger: split it.
MAX_WORKBOOK_ROWS = 20_000


@dataclass(frozen=True)
class Sheet:
    title: str          # the sheet's name in the workbook (Excel allows 31 characters)
    about: str          # one line for the Read me sheet
    viewset: str = ""   # dotted path of the viewset whose import this sheet runs
    service: str = ""   # or: a service-run sheet ('users', 'stock', 'on_order')
    perms: tuple = field(default_factory=tuple)  # a service sheet's permissions


_VS = "Tracker.viewsets."
SHEETS: tuple[Sheet, ...] = (
    Sheet("Users", "Everyone who will sign in. Loading invites each new one by email.",
          service="users", perms=("add_user", "change_user")),
    Sheet("Companies", "Customers and suppliers.", _VS + "core.CompanyViewSet"),
    Sheet("External Contacts", "People at those companies who get notifications.",
          _VS + "notifications.ExternalContactViewSet"),
    Sheet("Storage Locations", "Where stock is kept, if you keep a managed list.",
          _VS + "mes_standard.StorageLocationViewSet"),
    Sheet("Part Types", "Every part number you make, buy or repair.",
          _VS + "mes_lite.PartTypeViewSet"),
    Sheet("Materials", "Bulk materials and consumables bought by quantity.",
          _VS + "mes_standard.MaterialViewSet"),
    Sheet("BOMs", "Bill-of-material lines, one row per line. Each BOM loads as a draft.",
          _VS + "mes_standard.BOMLineViewSet"),
    Sheet("Equipment Types", "Kinds of machine and gauge.", _VS + "mes_lite.EquipmentTypeViewSet"),
    Sheet("Equipment", "Each machine and gauge.", _VS + "mes_lite.EquipmentViewSet"),
    Sheet("Tooling", "Fixtures and tools.", _VS + "scheduling.FixtureViewSet"),
    Sheet("Work Centers", "Groups of machines or people that work is scheduled on.",
          _VS + "mes_standard.WorkCenterViewSet"),
    Sheet("Shifts", "Shift patterns.", _VS + "mes_standard.ShiftViewSet"),
    Sheet("Plant Closures", "Holidays and shutdowns.",
          _VS + "scheduling.PlantCalendarExceptionViewSet"),
    Sheet("Changeovers", "Changeover times between part types on a work center.",
          _VS + "scheduling_setup.WorkCenterChangeoverViewSet"),
    Sheet("Step Timings", "Standard times per step (the steps must already exist).",
          _VS + "scheduling_setup.StepTimingViewSet"),
    Sheet("Machine Eligibility", "Which machines can run which step.",
          _VS + "scheduling_setup.StepEquipmentAffinityViewSet"),
    Sheet("Measurements", "What is measured at each step, with limits.",
          _VS + "qms.MeasurementsDefinitionViewSet"),
    Sheet("Sampling Rules", "Sampling rule sets and their rules.", _VS + "qms.SamplingRuleViewSet"),
    Sheet("Error Types", "The defect catalogue.", _VS + "qms.ErrorTypeViewSet"),
    Sheet("Repair Codes", "Reman repair codes.", _VS + "reman.RepairCodeViewSet"),
    Sheet("Life Limits", "Life-limit definitions (hours, cycles, calendar).",
          _VS + "life_tracking.LifeLimitDefinitionViewSet"),
    Sheet("Part Type Life Limits", "Which life limits apply to which part type.",
          _VS + "life_tracking.PartTypeLifeLimitViewSet"),
    Sheet("Order Milestones", "Milestone templates and their milestones.",
          _VS + "mes_lite.MilestoneViewSet"),
    Sheet("Job Roles", "Roles people hold.", _VS + "training.JobRoleViewSet"),
    Sheet("Training Types", "Trainings and qualifications.", _VS + "training.TrainingTypeViewSet"),
    Sheet("Training Requirements", "Which training each step, process, machine or role needs.",
          _VS + "training.TrainingRequirementViewSet"),
    Sheet("Training Records", "Who is trained in what, as of today.",
          _VS + "training.TrainingRecordViewSet"),
    Sheet("Calibration Records", "Each gauge's most recent calibration.",
          _VS + "calibration.CalibrationRecordViewSet"),
    Sheet("Orders", "Open customer orders.", _VS + "mes_lite.OrdersViewSet"),
    Sheet("Work Orders", "Open work orders (their process must already exist).",
          _VS + "mes_lite.WorkOrderViewSet"),
    Sheet("Parts", "Parts in work, with the step each is at now.", _VS + "mes_lite.PartsViewSet"),
    Sheet("Stock on Hand", "Lots on the shelf today, already accepted.",
          service="stock", perms=("add_materiallot", "change_materiallot")),
    Sheet("On Order", "Open purchase-order lines, expected in.",
          service="on_order", perms=("add_materiallot", "change_materiallot")),
)
SHEET_BY_TITLE = {s.title: s for s in SHEETS}


# ---------------------------------------------------------------------------
# Viewsets, as the per-table import uses them
# ---------------------------------------------------------------------------

def _viewset(sheet: Sheet, request):
    """The sheet's viewset, set up as for a request to its import action."""
    vs = import_string(sheet.viewset)()
    vs.request, vs.args, vs.kwargs = request, (), {}
    vs.format_kwarg, vs.action = None, "import_data"
    return vs


def _perms(sheet: Sheet, request) -> tuple[str, ...]:
    """What loading this sheet needs: add for new rows and change for upserts, as the
    per-table import requires."""
    if sheet.service:
        return sheet.perms
    name = _viewset(sheet, request)._get_model()._meta.model_name
    return (f"add_{name}", f"change_{name}")


def _missing_perms(sheet: Sheet, request) -> list[str]:
    return [p for p in _perms(sheet, request) if not request.user.has_tenant_perm(p)]


def _service_columns(sheet: Sheet) -> tuple[list[str], dict, dict]:
    """A service sheet's (columns, help per column, header→key map)."""
    if sheet.service == "stock":
        from Tracker.services.mes import stock_import as m
        return m.TEMPLATE_COLUMNS, m.COLUMN_HELP, m.FIELD_MAP
    if sheet.service == "on_order":
        from Tracker.services.mes import expected_receipt_import as m
        return m.TEMPLATE_COLUMNS, {}, m.FIELD_MAP
    # users: the bulk-users workbook's columns (services.core.user_reconcile)
    return (["Email*", "First Name", "Last Name", "Group", "Status"],
            {"Email*": "How they sign in. One row per person.",
             "Group": "The group(s) they belong to, e.g. Operator. Several: separate with ;",
             "Status": "Active or Inactive. Blank means Active for someone new."},
            {})


# ---------------------------------------------------------------------------
# The blank workbook
# ---------------------------------------------------------------------------

def build_template(request) -> bytes:
    """A blank workbook: a Read me sheet, then one sheet per table in load order, each
    with the columns its import accepts (hints on the header cells, dropdowns for
    fixed choices)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from Tracker.services.template_generator import HEADER_FILL, HEADER_FONT, THIN_BORDER

    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"

    viewsets = {s.title: _viewset(s, request) for s in SHEETS if s.viewset}
    sheet_for_model = {vs._get_model(): title for title, vs in viewsets.items()}

    for sheet in SHEETS:
        ws = wb.create_sheet(sheet.title)
        if sheet.viewset:
            generator = viewsets[sheet.title]._importable_generator()
            generator.write_data_sheet(ws, sheet_for_model)
            continue
        columns, help_text, _ = _service_columns(sheet)
        for col, name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, THIN_BORDER
            if name in help_text:
                cell.comment = Comment(help_text[name], "Import template")
            ws.column_dimensions[cell.column_letter].width = max(len(name), 14) + 2
        ws.freeze_panes = "A2"

    lines = [
        ("Master migration workbook", Font(size=14, bold=True)),
        ("Fill in the sheets you need and leave the rest empty. Upload it on the Data "
         "Management page: a dry run shows what every row would do, and nothing is kept "
         "until you load it.", None),
        ("", None),
        ("How it loads", Font(bold=True)),
        ("• The sheets load in the order below, so a row may name anything made on an "
         "earlier sheet — a part type from Part Types, a company from Companies.", None),
        ("• Nothing is loaded unless every row loads. Fix what the dry run reports and "
         "upload the workbook again.", None),
        ("• Uploading it again updates what it loaded before (matched by name, number or "
         "email); nothing is ever deleted.", None),
        ("• Columns marked * are required for a new row. Hover a column's header for what "
         "goes in it.", None),
        ("• Processes and steps aren't on this workbook: build them in the process editor "
         "first, then load step timings, measurements and work orders against them.", None),
        ("", None),
        ("Sheets, in load order", Font(bold=True)),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        cell = readme.cell(row=r, column=1, value=text)
        if font:
            cell.font = font
    r = len(lines) + 1
    for i, sheet in enumerate(SHEETS, start=1):
        readme.cell(row=r, column=1, value=f"{i}. {sheet.title}").font = Font(bold=True)
        readme.cell(row=r, column=2, value=sheet.about)
        r += 1
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 90

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Reading an uploaded workbook
# ---------------------------------------------------------------------------

def read_workbook(file, request) -> tuple[dict[str, list[dict]], list[str]]:
    """Each known sheet's rows, parsed as its import parses them, plus the names of
    sheets that aren't the workbook's (reported, not loaded). Raises ValueError for a
    file that can't be read or is too big."""
    import pandas as pd
    from Tracker.services.csv_utils import MAX_UPLOAD_BYTES, parse_excel_file

    name = getattr(file, "name", "") or ""
    if not name.lower().endswith(".xlsx"):
        raise ValueError("Upload the workbook as .xlsx.")
    size = getattr(file, "size", None)
    if size and size > MAX_UPLOAD_BYTES:
        raise ValueError(f"The file is {size // (1024 * 1024)} MB; the limit is "
                         f"{MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    data = file.read()
    try:
        names = pd.ExcelFile(io.BytesIO(data)).sheet_names
    except Exception:
        raise ValueError("That file can't be read as an Excel workbook.")

    rows_by_sheet, total = {}, 0
    for title in names:
        sheet = SHEET_BY_TITLE.get(title)
        if sheet is None:
            continue
        if sheet.viewset:
            field_map = getattr(_viewset(sheet, request), "csv_field_mapping", {}) or {}
        else:
            field_map = _service_columns(sheet)[2]
        rows, _ = parse_excel_file(io.BytesIO(data), field_map, sheet_name=title)
        rows = [r for r in rows if any(v not in (None, "") for v in r.values())]
        if rows:
            rows_by_sheet[title] = rows
            total += len(rows)
    if total > MAX_WORKBOOK_ROWS:
        raise ValueError(f"{total} rows is more than one workbook takes ({MAX_WORKBOOK_ROWS}). "
                         "Load it in two workbooks: setup first, then the open work.")
    ignored = [n for n in names if n not in SHEET_BY_TITLE and n != "Read me"]
    return rows_by_sheet, ignored


# ---------------------------------------------------------------------------
# Running it
# ---------------------------------------------------------------------------

def _text(errors) -> str:
    """A row's errors — a DRF detail of any shape — as one readable line."""
    if isinstance(errors, dict):
        return "; ".join(f"{k}: {_text(v)}" if k not in ("non_field_errors", "detail") else _text(v)
                         for k, v in errors.items())
    if isinstance(errors, (list, tuple)):
        return "; ".join(_text(e) for e in errors)
    return str(errors)


def _sheet_result(title, created=0, updated=0, unchanged=0, errors=0, rows=None, detail=""):
    return {"sheet": title, "created": created, "updated": updated, "unchanged": unchanged,
            "errors": errors, "rows": rows or [], "detail": detail}


# A data row's spreadsheet row number: the header is row 1.
def _excel_row(n: int) -> int:
    return n + 1


def _run_viewset_sheet(sheet, rows, request):
    from Tracker.serializers.csv_import import ImportMode
    vs = _viewset(sheet, request)
    body = vs._process_import_inline(rows, ImportMode.UPSERT, vs.get_csv_import_serializer(),
                                     request.tenant, request.user).data
    summary = body.get("summary", {})
    problems = []
    for r in body.get("results", []):
        if r.get("status") == "error":
            problems.append({"row": _excel_row(r["row"]), "outcome": "error",
                             "detail": _text(r.get("errors"))})
        elif r.get("warnings"):
            problems.append({"row": _excel_row(r["row"]), "outcome": "warning",
                             "detail": _text(r["warnings"])})
    return _sheet_result(sheet.title, summary.get("created", 0), summary.get("updated", 0),
                         0, summary.get("errors", 0), problems)


def _run_service_sheet(sheet, rows, request):
    tenant, user = request.tenant, request.user
    if sheet.service == "users":
        from Tracker.services.core.user_reconcile import reconcile_user_row
        counts = {"created": 0, "updated": 0, "unchanged": 0, "error": 0}
        problems = []
        for n, row in enumerate(rows, start=1):
            out = reconcile_user_row(row=row, tenant=tenant, acting_user=user)
            counts[out["outcome"]] = counts.get(out["outcome"], 0) + 1
            if out["outcome"] == "error":
                problems.append({"row": _excel_row(n), "outcome": "error", "detail": out["error"]})
            elif out.get("warnings"):
                problems.append({"row": _excel_row(n), "outcome": "warning",
                                 "detail": _text(out["warnings"])})
        return _sheet_result(sheet.title, counts["created"], counts["updated"],
                             counts["unchanged"], counts["error"], problems)

    if sheet.service == "stock":
        from Tracker.services.mes.stock_import import import_stock_rows as run
    else:
        from Tracker.services.mes.expected_receipt_import import import_expected_rows as run
    body = run(tenant=tenant, rows=rows)
    problems = []
    for r in body["rows"]:
        if r["outcome"] == "ERROR":
            problems.append({"row": _excel_row(r["row"]), "outcome": "error", "detail": r["detail"]})
        elif r.get("detail"):
            problems.append({"row": _excel_row(r["row"]), "outcome": "note", "detail": r["detail"]})
    return _sheet_result(sheet.title, body["created"] + body.get("reopened", 0), body["updated"],
                         body["unchanged"] + body.get("in_use", 0), body["errors"], problems)


def run_workbook(rows_by_sheet: dict[str, list[dict]], request, *, dry_run: bool,
                 on_sheet=None) -> dict:
    """Load the sheets in order, in one transaction. Rolled back for a dry run, and for
    a real load that hit any error. `request` needs `.user` and `.tenant` (a worker
    passes a stand-in); `on_sheet(i, n, title)` reports progress."""
    todo = [s for s in SHEETS if rows_by_sheet.get(s.title)]
    results = []
    with transaction.atomic():
        for i, sheet in enumerate(todo, start=1):
            if on_sheet:
                on_sheet(i, len(todo), sheet.title)
            rows = rows_by_sheet[sheet.title]
            missing = _missing_perms(sheet, request)
            if missing:
                results.append(_sheet_result(sheet.title, errors=len(rows), detail=(
                    f"You don't have permission to load this sheet ({', '.join(missing)}).")))
                continue
            # A savepoint per sheet: an error the sheet's own import didn't catch must
            # not leave the transaction aborted for every sheet after it.
            sid = transaction.savepoint()
            try:
                run = _run_service_sheet if sheet.service else _run_viewset_sheet
                results.append(run(sheet, rows, request))
                transaction.savepoint_commit(sid)
            except Exception as e:  # noqa: BLE001 — reported against the sheet
                transaction.savepoint_rollback(sid)
                results.append(_sheet_result(sheet.title, errors=len(rows),
                                             detail=f"The sheet couldn't be loaded: {e}"))
        errors = sum(r["errors"] for r in results)
        loaded = not dry_run and errors == 0
        if not loaded:
            transaction.set_rollback(True)
    return {
        "dry_run": dry_run,
        "loaded": loaded,
        "totals": {k: sum(r[k] for r in results)
                   for k in ("created", "updated", "unchanged", "errors")},
        "sheets": results,
    }
