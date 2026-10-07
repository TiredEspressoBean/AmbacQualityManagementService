"""What a gauge measured while it may have been out of tolerance (ISO 9001 7.1.5.2).

A failed calibration sets the gauge OUT_OF_SERVICE (`apply_calibration_result_to_equipment`),
which protects the next part. It says nothing about the parts already measured: a gauge
doesn't fail the moment it's checked, it drifts during the interval, so everything it
measured since its last good calibration is in question. This answers "which parts?".

It does not act on the answer. Re-inspecting, accepting on other evidence, telling a
customer or raising an NCR are judgements for whoever reads the list.

The window runs from the last good calibration (PASS, or LIMITED — restricted use, not
unfit) to the check that found the gauge unfit. Both ends are taken as whole plant-local
days: a calibration record has a date, not a time, so a report filed that day may be
either side of the check, and a suspect list errs toward including it.

"Found unfit" is a FAIL, or any result with `as_found_in_tolerance = False` — a gauge
found out of tolerance, adjusted and passed is a PASS that measured wrongly until then.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta

GOOD_RESULTS = ('PASS', 'LIMITED')


class NotFoundUnfit(ValueError):
    """The calibration found the gauge fit for use — nothing it measured is in question."""


@dataclass(frozen=True)
class ExposedReport:
    id: str
    report_number: str
    part: str            # the part's ERP id, or the material lot's number (receiving inspection)
    part_id: str | None
    material_lot_id: str | None
    step: str
    created_at: datetime
    inspector: str
    status: str          # PASS / FAIL / PENDING
    role: str            # how the gauge was attached to the report (GAUGE, normally)


@dataclass(frozen=True)
class CalibrationExposure:
    calibration_id: str
    equipment_id: str
    equipment_name: str
    result: str
    as_found_in_tolerance: bool | None
    window_start: date | None        # None: no earlier good calibration on record
    window_start_calibration_id: str | None
    window_start_result: str | None
    window_end: date
    reports: tuple[ExposedReport, ...]


def is_found_unfit(record) -> bool:
    return record.result == 'FAIL' or record.as_found_in_tolerance is False


def exposure_for(record) -> CalibrationExposure:
    """The quality reports the record's gauge was used on between its last good
    calibration and this one. Raises NotFoundUnfit for a calibration that found the gauge
    fit for use."""
    from Tracker.models import CalibrationRecord, QualityReportEquipment, QualityReports
    from Tracker.services.core.clock import plant_tz

    if not is_found_unfit(record):
        raise NotFoundUnfit(
            "This calibration found the gauge in tolerance, so nothing it measured is in question.")

    equipment = record.equipment
    # Every version of the gauge: a new version is the same instrument, so its history
    # (calibrations and the reports it was used on) spans the chain.
    chain = [v.pk for v in equipment.get_version_history()]

    last_good = (CalibrationRecord.objects  # tenant-safe: .objects auto-scopes
                 .filter(equipment_id__in=chain, archived=False, result__in=GOOD_RESULTS,
                         calibration_date__lt=record.calibration_date)
                 .exclude(pk=record.pk)
                 .order_by('-calibration_date', '-created_at')
                 .first())

    tz = plant_tz(record.tenant)
    end = datetime.combine(record.calibration_date + timedelta(days=1), time.min, tzinfo=tz)
    reports = (QualityReports.objects  # tenant-safe: .objects auto-scopes
               .filter(archived=False, equipment_links__equipment_id__in=chain,
                       created_at__lt=end)
               .select_related('part', 'material_lot', 'step', 'detected_by')
               .order_by('created_at'))
    if last_good:
        start = datetime.combine(last_good.calibration_date, time.min, tzinfo=tz)
        reports = reports.filter(created_at__gte=start)

    roles = {}
    for qr_id, role in (QualityReportEquipment.objects
                        .filter(equipment_id__in=chain, quality_report__in=reports)
                        .values_list('quality_report_id', 'role')):
        roles.setdefault(qr_id, role)

    rows = []
    seen = set()
    for qr in reports:
        if qr.pk in seen:  # one report can link the gauge twice (two roles)
            continue
        seen.add(qr.pk)
        if qr.part_id:
            part = qr.part.ERP_id
        elif qr.material_lot_id:
            part = qr.material_lot.lot_number
        else:
            part = ''
        rows.append(ExposedReport(
            id=str(qr.pk),
            report_number=qr.report_number or '',
            part=part,
            part_id=str(qr.part_id) if qr.part_id else None,
            material_lot_id=str(qr.material_lot_id) if qr.material_lot_id else None,
            step=qr.step.name if qr.step_id else '',
            created_at=qr.created_at,
            inspector=(qr.detected_by.get_full_name() or qr.detected_by.email) if qr.detected_by_id else '',
            status=qr.status,
            role=roles.get(qr.pk, ''),
        ))

    return CalibrationExposure(
        calibration_id=str(record.pk),
        equipment_id=str(equipment.pk),
        equipment_name=equipment.name,
        result=record.result,
        as_found_in_tolerance=record.as_found_in_tolerance,
        window_start=last_good.calibration_date if last_good else None,
        window_start_calibration_id=str(last_good.pk) if last_good else None,
        window_start_result=last_good.result if last_good else None,
        window_end=record.calibration_date,
        reports=tuple(rows),
    )


COLUMNS = ["Report", "Part / Lot", "Step", "Filed", "Inspector", "Result", "Gauge Role"]


def exposure_workbook(exposure: CalibrationExposure, tenant) -> bytes:
    """The list as an .xlsx, with the window it was drawn from — the window is the finding."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from Tracker.services.core.clock import plant_tz
    from Tracker.services.spreadsheet_safety import write_cell
    from Tracker.services.template_generator import HEADER_FILL, HEADER_FONT

    tz = plant_tz(tenant)
    wb = Workbook()
    ws = wb.active
    ws.title = "Measured"
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[cell.column_letter].width = max(len(name), 14) + 2
    for r, row in enumerate(exposure.reports, start=2):
        values = [row.report_number, row.part, row.step,
                  row.created_at.astimezone(tz).strftime("%Y-%m-%d %H:%M"),
                  row.inspector, row.status, row.role]
        for c, value in enumerate(values, start=1):
            write_cell(ws, r, c, value)
    ws.freeze_panes = "A2"

    about = wb.create_sheet("About")
    lines = [(f"What {exposure.equipment_name} measured", Font(size=13, bold=True)),
             (window_sentence(exposure), None),
             (f"{len(exposure.reports)} quality report(s).", None),
             ("ISO 9001 7.1.5.2: a gauge found unfit may have been out of tolerance since its "
              "last good calibration, so these measurements are in question. Deciding what to "
              "do — re-inspect, accept on other evidence, notify a customer, raise an NCR — is "
              "not done by UQMES.", None)]
    for i, (text, font) in enumerate(lines, start=1):
        cell = about.cell(row=i, column=1, value=text)
        if font:
            cell.font = font
    about.column_dimensions["A"].width = 110
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def window_sentence(exposure: CalibrationExposure) -> str:
    found = ("failed check" if exposure.result == 'FAIL'
             else "check that found it out of tolerance")
    if exposure.window_start is None:
        return (f"No earlier passing calibration on record: everything measured up to "
                f"{exposure.window_end:%Y-%m-%d} ({found}).")
    last = "last limited-use calibration" if exposure.window_start_result == 'LIMITED' \
        else "last passing calibration"
    return (f"Measured between {exposure.window_start:%Y-%m-%d} ({last}) and "
            f"{exposure.window_end:%Y-%m-%d} ({found}).")
